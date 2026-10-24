"""Writes yzu_reference.docx, the style template of the YZU-layout Word copy (YZU thesis format guidelines, G-en II.5 to II.8 and III.4):
A4; margins top 3.5 cm, left 4 cm, right 2 cm, bottom 2 cm; page number centred 1 cm from the bottom; English text at 1.5 line spacing;
chapter title centred at the top of a new page with double space under it; Times New Roman 12 pt (the font of YZU's own templates; the guideline sets no font).
Run: python3 make_yzu_reference.py (needs pandoc and python-docx)."""
import subprocess, docx
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
subprocess.run(["pandoc", "-o", "/tmp/ref_default.docx", "--print-default-data-file", "reference.docx"], check=True)
d = docx.Document("/tmp/ref_default.docx")
by = {s.name.lower(): s for s in d.styles}
sec = d.sections[0]
sec.page_height, sec.page_width = Mm(297), Mm(210)
sec.top_margin, sec.left_margin, sec.right_margin, sec.bottom_margin = Mm(35), Mm(40), Mm(20), Mm(20)
sec.footer_distance = Mm(10); sec.header_distance = Mm(15)
TEXT_W = 150.0
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
def setfont(st, size, bold=None, italic=None):
    st.font.name = FONT
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    st.font.size = Pt(size); st.font.color.rgb = BLACK
    if bold is not None: st.font.bold = bold
    if italic is not None: st.font.italic = italic
def para(st, rule, line, before=0, after=0, align=None, keep=None, pbb=None):
    pf = st.paragraph_format
    pf.line_spacing_rule = rule
    pf.line_spacing = line
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if align is not None: pf.alignment = align
    if keep: pf.keep_with_next = True
    if pbb: pf.page_break_before = True
J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
M, X = WD_LINE_SPACING.MULTIPLE, WD_LINE_SPACING.EXACTLY
def get_or_make(name):
    if name.lower() in by: return by[name.lower()]
    st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH); st.base_style = by["normal"]; by[name.lower()] = st; return st
for n in ("normal", "body text", "first paragraph"):
    setfont(by[n], 12); para(by[n], M, 1.5, 0, 6, J)
setfont(by["compact"], 12); para(by["compact"], M, 1.15, 0, 3, L)
setfont(by["block text"], 12); para(by["block text"], M, 1.5, 3, 6, L)
by["block text"].paragraph_format.left_indent = Mm(10); by["block text"].paragraph_format.right_indent = Mm(10)
setfont(by["title"], 18, True); para(by["title"], M, 1.3, 90, 18, C)
setfont(by["subtitle"], 12, False, True); para(by["subtitle"], M, 1.3, 0, 40, C)
for n in ("author", "date"): setfont(by[n], 12); para(by[n], M, 1.3, 0, 4, C)
setfont(by["heading 1"], 16, True, False); para(by["heading 1"], M, 1.3, 0, 24, C, True, True)
setfont(by["heading 2"], 13, True, False); para(by["heading 2"], M, 1.3, 14, 6, L, True)
setfont(by["heading 3"], 12, True, True); para(by["heading 3"], M, 1.3, 10, 4, L, True)
setfont(by["heading 4"], 12, True, False); para(by["heading 4"], M, 1.3, 8, 3, L, True)
setfont(by["footnote text"], 10); para(by["footnote text"], M, 1.0, 0, 2, L)
fh = get_or_make("Front Heading"); setfont(fh, 16, True, False); para(fh, M, 1.3, 0, 24, C, True, True)
fs = get_or_make("Field Slot"); setfont(fs, 12); para(fs, M, 1.0, 0, 0, L)
tc = get_or_make("Table Caption"); setfont(tc, 11, False, False); para(tc, M, 1.15, 12, 4, L, True)
fc = get_or_make("Figure Caption"); setfont(fc, 11, False, False); para(fc, M, 1.15, 4, 3, L)
tn = get_or_make("Table Note"); setfont(tn, 9.5, False, False); para(tn, M, 1.1, 2, 12, J)
tn.paragraph_format.alignment = L
eq = get_or_make("Equation"); setfont(eq, 12); para(eq, M, 1.3, 6, 8, L); eq.paragraph_format.keep_together = True
eq.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER); eq.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
for name, left, bold in (("toc 1", 0, True), ("toc 2", 7, False), ("table of figures", 0, False)):
    st = get_or_make(name); setfont(st, 12, bold); para(st, M, 1.15, 4 if name == "toc 1" else 0, 2, L)
    st.paragraph_format.left_indent = Mm(left)
    st.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph(); p.alignment = C
r = p.add_run()
for t, txt in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
    if t == "instr": el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    elif t == "text": el = OxmlElement("w:t"); el.text = txt
    else: el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
    r._r.append(el)
r.font.size = Pt(11); r.font.name = FONT
d.save("yzu_reference.docx"); print("yzu_reference.docx written")
