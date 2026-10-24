"""Writes docx_reference.docx, the style template of the condensed Word copy: A4, Times New Roman 10.5 pt, 1.08 spacing, centred page number,
an "Equation" paragraph style with a centre and a right tab. Run: python3 make_docx_reference.py (needs pandoc and python-docx)."""
import subprocess, docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
subprocess.run(["pandoc", "-o", "/tmp/ref_default.docx", "--print-default-data-file", "reference.docx"], check=True)
d = docx.Document("/tmp/ref_default.docx")
by = {s.name.lower(): s for s in d.styles}
get = by.get
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
TEXT_W = 21.0 - 4.4
FONT = "Times New Roman"
def setfont(st, size, bold=None, italic=None):
    st.font.name = FONT
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    st.font.size = Pt(size); st.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None: st.font.bold = bold
    if italic is not None: st.font.italic = italic
spec = {"normal": (10.5, None, None), "body text": (10.5, None, None), "first paragraph": (10.5, None, None), "compact": (10.5, None, None),
        "block text": (10.5, None, None), "title": (17, True, None), "subtitle": (11.5, False, True), "author": (10.5, None, None), "date": (10.5, None, None),
        "heading 1": (15, True, False), "heading 2": (12.5, True, False), "heading 3": (11, True, False), "heading 4": (10.5, True, False),
        "table caption": (10, False, False), "image caption": (10, False, False), "caption": (10, False, False), "footnote text": (9.5, None, None)}
for n, (sz, b, i) in spec.items():
    if get(n) is not None: setfont(get(n), sz, b, i)
for n in ("normal", "body text", "first paragraph"):
    pf = get(n).paragraph_format
    pf.line_spacing = 1.08; pf.space_after = Pt(4); pf.space_before = Pt(0); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
get("compact").paragraph_format.space_after = Pt(1); get("compact").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
get("block text").paragraph_format.left_indent = Cm(0.8); get("block text").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
for n, b, a in (("heading 1", 14, 6), ("heading 2", 10, 3), ("heading 3", 6, 2)):
    pf = get(n).paragraph_format; pf.space_before = Pt(b); pf.space_after = Pt(a); pf.keep_with_next = True
for n in ("title", "subtitle", "author", "date"): get(n).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
if "equation" not in by:
    eq = d.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH); eq.base_style = get("normal")
else: eq = by["equation"]
setfont(eq, 10.5)
pf = eq.paragraph_format; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.space_before = Pt(4); pf.space_after = Pt(6); pf.keep_together = True
pf.tab_stops.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER); pf.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run()
for t, txt in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
    if t == "instr": el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    elif t == "text": el = OxmlElement("w:t"); el.text = txt
    else: el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
    r._r.append(el)
r.font.size = Pt(9); r.font.name = FONT
d.save("docx_reference.docx"); print("docx_reference.docx written")
