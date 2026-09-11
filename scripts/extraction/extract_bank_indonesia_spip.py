#!/usr/bin/env python3
"""Extract Bank Indonesia SPIP payment tables without changing source values.

The source workbook presents years as repeated blocks containing a published
annual column followed by monthly columns.  This script preserves those two
frequencies separately and records the source sheet/row for every observation.
It does not interpolate missing values or turn the partial 2025 monthly block
into a full-year observation.
"""

from __future__ import annotations

import csv
import hashlib
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "sources/bank_indonesia_payments/raw/SPIP-Desember-2025.xlsx"
OUT = ROOT / "data/payments"
SOURCE_LABEL = "Bank Indonesia SPIP December 2025 workbook"

NATIONAL_SHEETS = ("5a", "5c", "5e", "5g", "6", "7")
REGIONAL_SHEETS = ("5b", "5d", "5f", "5h")
MONTHS = {
    "Jan": 1,
    "Feb": 2,
    "Mar": 3,
    "Apr": 4,
    "May": 5,
    "Jun": 6,
    "Jul": 7,
    "Aug": 8,
    "Sep": 9,
    "Oct": 10,
    "Nov": 11,
    "Dec": 12,
}


def numeric(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def time_columns(ws) -> tuple[list[tuple[int, int]], list[tuple[int, int, int]]]:
    annual: list[tuple[int, int]] = []
    monthly: list[tuple[int, int, int]] = []
    current_year: int | None = None
    for col in range(4, ws.max_column + 1):
        year_cell = ws.cell(5, col).value
        month_cell = ws.cell(6, col).value
        if isinstance(year_cell, (int, float)):
            current_year = int(year_cell)
        elif isinstance(year_cell, str) and year_cell.strip().isdigit():
            current_year = int(year_cell.strip())
        if current_year is None:
            continue
        month = MONTHS.get(str(month_cell).strip()) if month_cell is not None else None
        if month:
            monthly.append((col, current_year, month))
        elif year_cell is not None and str(year_cell).strip() not in {"", "COMPONENTS"}:
            annual.append((col, current_year))
    return annual, monthly


def sheet_title(ws) -> str:
    return str(ws.cell(2, 1).value or ws.cell(1, 1).value or ws.title).strip()


def observation_rows(ws) -> Iterable[int]:
    for row in range(7, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if label is None or str(label).strip() == "":
            continue
        yield row


def extract_monthly(wb, sheets: Iterable[str]) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for sheet in sheets:
        ws = wb[sheet]
        _, columns = time_columns(ws)
        for row in observation_rows(ws):
            label = str(ws.cell(row, 2).value).strip()
            unit = ws.cell(row, 3).value
            indicator_id = ws.cell(row, 1).value
            for col, year, month in columns:
                value = ws.cell(row, col).value
                if not numeric(value):
                    continue
                records.append(
                    {
                        "source_sheet": sheet,
                        "source_table": sheet_title(ws),
                        "source_row": row,
                        "indicator_id": indicator_id,
                        "indicator_label": label,
                        "unit": unit,
                        "year": year,
                        "month": month,
                        "period": f"{year:04d}-{month:02d}",
                        "value_as_published": value,
                        "source_file": SOURCE.name,
                    }
                )
    return records


def extract_annual(wb, sheets: Iterable[str]) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for sheet in sheets:
        ws = wb[sheet]
        columns, _ = time_columns(ws)
        for row in observation_rows(ws):
            label = str(ws.cell(row, 2).value).strip()
            unit = ws.cell(row, 3).value
            indicator_id = ws.cell(row, 1).value
            for col, year in columns:
                value = ws.cell(row, col).value
                if not numeric(value):
                    continue
                records.append(
                    {
                        "source_sheet": sheet,
                        "source_table": sheet_title(ws),
                        "source_row": row,
                        "indicator_id": indicator_id,
                        "indicator_label": label,
                        "unit": unit,
                        "year": year,
                        "value_as_published": value,
                        "source_file": SOURCE.name,
                    }
                )
    return records


CORE_ROWS = {
    ("5a", 13): "atm_debit_shopping_volume",
    ("5a", 19): "atm_debit_shopping_value",
    ("5c", 13): "credit_card_shopping_volume",
    ("5c", 21): "credit_card_shopping_value",
    ("5e", 7): "electronic_money_instruments",
    ("5e", 16): "electronic_money_total_volume",
    ("5e", 17): "electronic_money_shopping_volume",
    ("5e", 26): "electronic_money_total_value",
    ("5e", 27): "electronic_money_shopping_value",
    ("5g", 17): "apmk_electronic_money_merchants",
    ("6", 11): "nonbank_domestic_transfer_volume",
    ("6", 16): "nonbank_domestic_transfer_value",
    ("7", 8): "proprietary_channel_total_volume",
    ("7", 9): "proprietary_channel_total_value",
    ("7", 22): "mobile_banking_total_volume",
    ("7", 23): "mobile_banking_payment_purchase_volume",
    ("7", 26): "mobile_banking_total_value",
    ("7", 27): "mobile_banking_payment_purchase_value",
    ("7", 32): "internet_banking_total_volume",
    ("7", 33): "internet_banking_payment_purchase_volume",
    ("7", 36): "internet_banking_total_value",
    ("7", 37): "internet_banking_payment_purchase_value",
}


def core(records: list[dict[str, object]]) -> list[dict[str, object]]:
    selected: list[dict[str, object]] = []
    for record in records:
        key = (str(record["source_sheet"]), int(record["source_row"]))
        if key not in CORE_ROWS:
            continue
        selected.append({"series_id": CORE_ROWS[key], **record})
    return selected


def write_csv(path: Path, records: list[dict[str, object]]) -> None:
    if not records:
        raise ValueError(f"No records for {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)


def write_inventory(paths: list[Path], workbook_hash: str) -> None:
    rows = []
    for path in paths:
        with path.open(encoding="utf-8") as handle:
            count = sum(1 for _ in handle) - 1
        rows.append(
            {
                "file": str(path.relative_to(ROOT)),
                "data_rows": count,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "source_workbook_sha256": workbook_hash,
                "status": "mechanical_extract_from_official_workbook",
            }
        )
    write_csv(OUT / "bank_indonesia_spip_extract_inventory_2026-09-11.csv", rows)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    workbook_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    # The workbook is small enough for normal mode; random cell access in
    # openpyxl's read-only mode repeatedly reparses XML and is much slower.
    wb = load_workbook(SOURCE, read_only=False, data_only=True)

    national_monthly = extract_monthly(wb, NATIONAL_SHEETS)
    regional_monthly = extract_monthly(wb, REGIONAL_SHEETS)
    national_annual = extract_annual(wb, NATIONAL_SHEETS)
    core_monthly = core(national_monthly)
    core_annual = core(national_annual)

    paths = [
        OUT / "bank_indonesia_spip_national_monthly_2009_2025.csv",
        OUT / "bank_indonesia_spip_regional_monthly_2009_2025.csv",
        OUT / "bank_indonesia_spip_national_published_annual_2009_2024.csv",
        OUT / "bank_indonesia_payments_core_monthly_2009_2025.csv",
        OUT / "bank_indonesia_payments_core_published_annual_2009_2024.csv",
    ]
    for path, records in zip(
        paths,
        [national_monthly, regional_monthly, national_annual, core_monthly, core_annual],
        strict=True,
    ):
        write_csv(path, records)
    write_inventory(paths, workbook_hash)

    counts = Counter(r["source_sheet"] for r in national_monthly)
    print(f"source_sha256={workbook_hash}")
    for path in paths:
        with path.open(encoding="utf-8") as handle:
            print(f"{path.relative_to(ROOT)} rows={sum(1 for _ in handle) - 1}")
    print("national_monthly_by_sheet", dict(sorted(counts.items())))


if __name__ == "__main__":
    main()
