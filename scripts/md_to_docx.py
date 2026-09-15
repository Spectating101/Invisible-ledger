# -*- coding: utf-8 -*-
"""Minimal markdown -> DOCX for the thesis manuscript draft.

Deliberately plain: this is a working draft, not the YZU thesis package. Final
production formatting (front matter, TOC, list of tables, official margins) belongs
to the stage after the sample boundary is frozen.

    python scripts/md_to_docx.py in.md out.docx
"""
import sys, re, io
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, out = sys.argv[1], sys.argv[2]
md = io.open(src, encoding='utf8').read()
d = Document()
for s in d.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.5); s.left_margin = s.right_margin = Cm(2.5)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)

def inline(p, text):
    for part in re.split(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'): r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('*') and part.endswith('*'):  r=p.add_run(part[1:-1]); r.italic=True
        elif part.startswith('`') and part.endswith('`'):  r=p.add_run(part[1:-1]); r.font.name='Courier New'
        else: p.add_run(part)

lines = md.split('\n'); i = 0
while i < len(lines):
    l = lines[i].rstrip()
    if l.startswith('|'):                                   # table
        block=[]
        while i < len(lines) and lines[i].lstrip().startswith('|'):
            row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(set(c) <= set('-: ') for c in row): block.append(row)
            i += 1
        if block:
            t=d.add_table(rows=0, cols=len(block[0])); t.style='Table Grid'
            for ri,row in enumerate(block):
                tr=t.add_row()
                for j,v in enumerate(row[:len(block[0])]):
                    cell=tr.cells[j]; cell.text=''
                    p=cell.paragraphs[0]; inline(p, v)
                    for r in p.runs: r.font.size=Pt(9)
                    if ri==0:
                        for r in p.runs: r.bold=True
                    p.paragraph_format.space_after=Pt(0)
            d.add_paragraph().paragraph_format.space_after=Pt(4)
        continue
    if l.startswith('# '):    p=d.add_heading(l[2:].strip(), 0)
    elif l.startswith('### '):p=d.add_heading(l[4:].strip(), 2)
    elif l.startswith('## '): p=d.add_heading(l[3:].strip(), 1)
    elif l.startswith('> '):
        p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(1); inline(p, l[2:])
    elif l.startswith('- '):
        p=d.add_paragraph(style='List Bullet'); inline(p, l[2:])
    elif l.strip()=='---': pass
    elif l.strip():
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after=Pt(6); inline(p, l)
    i += 1
d.save(out)
print('wrote', out)
