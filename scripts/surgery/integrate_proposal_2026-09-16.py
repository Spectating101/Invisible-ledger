#!/usr/bin/env python3
"""Integrate review fixes into the author-approved 16 September proposal freeze.

Input is the approved DOCX (commit f51331d). Chapters 1-2 are the author's text and are edited
only by the exact replacements listed in CH12_EDITS. The abstract is completed. Chapters 3-8 and
Appendix A reproduce the frozen text (Invisible_Ledger_Thesis_Proposal_FROZEN_2026-09-16.pdf)
with the targeted edits marked "# EDIT" below. This is a proposal: no results chapter is added.

Usage: integrate_proposal_2026-09-16.py APPROVED.docx OUT.docx
Then run scripts/finalize_proposal_format.py OUT.docx
"""
import copy
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC, OUT = sys.argv[1], sys.argv[2]
doc = Document(SRC)
FONT = "Times New Roman"
J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER


FOOTNOTES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
%s</w:footnotes>"""

FOOTNOTE_BODY = """<w:footnote w:id="%d"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr>
<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman"/><w:vertAlign w:val="superscript"/><w:sz w:val="20"/></w:rPr><w:footnoteRef/></w:r>
<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve"> %s</w:t></w:r></w:p></w:footnote>"""


def add_footnotes(document, texts):
    """Attach a real footnotes part; python-docx has no API for it."""
    from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    body = "\n".join(FOOTNOTE_BODY % (i + 2, t) for i, t in enumerate(texts))
    part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES,
                (FOOTNOTES_XML % body).encode("utf-8"), document.part.package)
    document.part.relate_to(part, RT.FOOTNOTES)


def footnote_ref(paragraph, note_id):
    r = paragraph.add_run()
    font(r, 12)
    r._element.get_or_add_rPr().append(OxmlElement("w:vertAlign"))
    r._element.rPr.find(qn("w:vertAlign")).set(qn("w:val"), "superscript")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(note_id))
    r._element.append(ref)


def find(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise RuntimeError("paragraph not found: " + text)


def clear_between(start_el, end_el):
    node = start_el.getnext()
    while node is not None and node is not end_el:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def font(run, size, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_runs(p, text, size=12, bold=False, italic=False):
    for i, part in enumerate(re.split(r"_\{(\w+)\}", text)):   # "_{it}" marks a subscript
        if part:
            r = p.add_run(part)
            font(r, size, bold, italic)
            if i % 2:
                r.font.subscript = True


def replace_in(p, old, new):
    if old not in p.text:
        raise RuntimeError("text not found: " + old)
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    raise RuntimeError("text split across runs: " + old)


class Writer:
    after_table = False

    def __init__(self, anchor_el):
        self.anchor = anchor_el

    def _p(self):
        el = OxmlElement("w:p")
        self.anchor.addprevious(el)
        p = Paragraph(el, doc._body)
        if self.after_table:
            p.paragraph_format.space_before = Pt(6)
            self.after_table = False
        return p

    def h1(self, text):
        p = self._p(); p.style = "Heading 1"; p.alignment = L; add_runs(p, text, bold=True)

    def h2(self, text):
        p = self._p(); p.style = "Heading 2"; p.alignment = L; add_runs(p, text, bold=True)

    def para(self, text, align=J, bold=False, italic=False):
        p = self._p(); p.style = "Normal"; p.alignment = align
        add_runs(p, text, bold=bold, italic=italic)
        return p

    def hyp(self, text):
        p = self.para(text, L, bold=True)
        p.paragraph_format.keep_with_next = True

    def eq(self, text):
        p = self.para(text, C, italic=True)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True

    def caption(self, text, center=False):
        p = self._p(); p.style = "Normal"; p.alignment = C if center else L
        add_runs(p, text, size=9.5, bold=True, italic=True)
        p.paragraph_format.keep_with_next = True

    def element(self, el):
        self.anchor.addprevious(el)

    def table(self, rows):
        ncols = len(rows[0])
        tbl = OxmlElement("w:tbl")
        tblPr = OxmlElement("w:tblPr")
        for tag, attrs in [("w:tblStyle", {"w:val": "TableGrid"}), ("w:tblW", {"w:type": "auto", "w:w": "0"}),
                           ("w:jc", {"w:val": "center"}), ("w:tblLayout", {"w:type": "fixed"})]:
            e = OxmlElement(tag)
            for k, v in attrs.items():
                e.set(qn(k), v)
            tblPr.append(e)
        tbl.append(tblPr)
        grid = OxmlElement("w:tblGrid")
        for _ in range(ncols):
            g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(9547 // ncols)); grid.append(g)
        tbl.append(grid)
        for ri, row in enumerate(rows):
            tr = OxmlElement("w:tr")
            trPr = OxmlElement("w:trPr")
            if ri == 0:
                h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
            trPr.append(OxmlElement("w:cantSplit"))
            tr.append(trPr)
            for text in row:
                tc = OxmlElement("w:tc")
                tcPr = OxmlElement("w:tcPr")
                va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
                tc.append(tcPr)
                pel = OxmlElement("w:p"); tc.append(pel)
                add_runs(Paragraph(pel, None), text, size=9, bold=(ri == 0))
                tr.append(tc)
            tbl.append(tr)
        self.anchor.addprevious(tbl)
        self.after_table = True
        return Table(tbl, doc._body)


figure_el = copy.deepcopy(next(p._p for p in doc.paragraphs if p._p.xpath(".//w:drawing")))
ch2_h = find("2. Literature Review and Research Gap")._p
ch3_h = next(p._p for p in doc.paragraphs if p.text.strip().startswith("3. Theoretical Framework"))
refs_h = find("References")._p

# ------------------------------------------------------------------ abstract (completed)
abstract = Paragraph(find("Abstract")._p.getnext(), doc._body)
for r in abstract.runs[1:]:
    r._element.getparent().remove(r._element)
abstract.runs[0].text = (
    "Digital economy is measured through platform revenue, national e-commerce statistics and payment data, "
    "which are often used interchangeably even though each records a different part of the same activity. This "
    "study asks how much economic activity remains invisible when digital platforms are measured through their "
    "reported revenue. I define the invisible wedge as the difference between transaction value and "
    "platform-recognized revenue, measure it across Indonesian platform histories from 2020 to 2025, and compare "
    "the result with BPS-Statistics Indonesia e-commerce data and Bank Indonesia payment records. In 2023, three "
    "major platforms processed US$43.23 billion of transaction value against US$3.16 billion of recognized "
    "revenue, a ratio of 13.7 to 1. Across twelve "
    "year-to-year comparisons, revenue grew faster than transaction value in ten, so the relationship is not "
    "stable. National statistics place 98.5 percent of recent e-commerce growth outside the marketplace "
    "component, and payment measures grew up to eleven times as fast as e-commerce value. Eight platform "
    "businesses outside Indonesia provide external corroboration. The proposal treats these records as distinct "
    "measurement layers and tests whether substituting one for another changes conclusions about the scale and "
    "growth of digital activity.")
abstract.alignment = J

# ------------------------------------------------------------------ Chapters 1-2: author's text, exact edits only
CH12_EDITS = [
    ("This paper argues that doing so can produce systematically wrong",
     "This paper argues that doing so produces systematically wrong"),
    ("they may represent merchant receipts", "they represent merchant receipts"),
    ("Indonesia provides a useful empirical setting because",
     "Indonesia, the largest digital market in Southeast Asia, provides a clear empirical setting because"),
    ("with a median divergence of approximately 42 percentage points;",
     "with a median absolute gap of approximately 42 percentage points between the two growth rates;"),
    ("while selected third-party net revenue rose 53.2 percent", "while its third-party net revenue rose 53.2 percent"),
    ("The OECD Model Rules and the European Union's DAC7 regime apply",
     "The OECD Model Rules (OECD 2020) and the European Union's DAC7 regime (European Union 2021) apply"),
]
for old, new in CH12_EDITS:
    hits = [p for p in doc.paragraphs if old in p.text]
    if len(hits) != 1:
        raise RuntimeError(f"{len(hits)} matches for: {old}")
    replace_in(hits[0], old, new)

node = find("1. Introduction")._p.getnext()   # justify body paragraphs of Chapters 1-2
while node is not None and node is not ch3_h:
    if node.tag == qn("w:p"):
        par = Paragraph(node, doc._body)
        if par.style.name == "Normal" and par.text.strip() and par.alignment is None:
            par.alignment = J
    node = node.getnext()

# ------------------------------------------------------------------ Chapters 3-8 and Appendix A (frozen text)
clear_between(ch3_h.getprevious(), refs_h)
w = Writer(refs_h)

w.h1("3. Theoretical Framework and Hypotheses")
w.para("Digital platform activity can be observed through several records, but those records do not measure the "
       "same economic object. The framework begins with the difference between the value a platform processes and "
       "the revenue it recognizes, then asks whether that relationship remains stable over time and whether other "
       "records produce the same account of digital activity.")
w.para("For the platform analysis, transaction value and revenue are paired only when they refer to the same period "
       "and business activity. Broader-scope, reconstructed, "
       "national-statistical, and payment evidence is kept in separate analytical roles rather than combined into "
       "one sample.")
w.h2("3.1 Measurement Framework")
w.para("Let V_{it} denote transaction value processed by platform i in period t, and R_{it} the revenue recognized "
       "over the same matched period and business scope. The absolute invisible wedge is:")
w.eq("W_{it} = V_{it} − R_{it}")
w.para("W measures transaction value outside the platform's recognized revenue. It does not represent merchant "
       "profit, worker income, taxable income, unpaid tax, tax evasion, or value added.")
w.para("Because platforms differ in size, I also express the wedge relative to revenue. The Ecosystem Ratio is:")
w.eq("E_{it} = (V_{it} − R_{it}) / R_{it}")
w.para("An E of 9 means that for every one unit of platform revenue, nine additional units of transaction value lie "
       "outside the revenue line. The Ecosystem Ratio is constructed for this study rather than adopted from prior "
       "work. Its basis is the platform-economics result that booked revenue need not track the transaction value a "
       "multisided platform coordinates (Rochet and Tirole 2003; Hagiu and Wright 2015; Evans and Schmalensee 2016), "
       "together with the accounting result that the share of transaction value entering revenue is itself a "
       "reporting choice (International Accounting Standards Board 2014; De Franco, Kothari and Verdi 2011).")
w.para("A single ratio does not show whether transaction value and revenue move together through time. I therefore "
       "define annual growth divergence as:")
w.eq("D_{it} = g(V_{it}) − g(R_{it})")
w.para("A positive D means transaction value grew faster than revenue; a negative D means revenue grew faster. "
       "Together, W, E, and D measure the size of the transaction-revenue boundary and how that boundary changes "
       "over time.")
w.h2("3.2 Hypotheses")
w.para("The four hypotheses examine different layers of the same measurement problem. H1 tests the "
       "transaction-revenue boundary within platforms. H2 and H3 examine where activity and records outside "
       "platform revenue appear in national e-commerce evidence. H4 asks whether the different records ultimately "
       "produce different conclusions about digital economic activity. The evidence layers are analyzed "
       "separately, and none of the hypotheses is interpreted as a causal claim. Table 1 pairs each hypothesis "
       "with the evidence that tests it.")
# EDIT: H1 and H4 stated without "can", so that the tests can fail; "series" label removed
w.hyp("H1. Transaction value and platform revenue diverge materially within the same platform over time.")
w.para("If revenue were a stable proxy for the commerce processed by a platform, transaction value and revenue would "
       "move proportionally.")
w.hyp("H2. Aggregate e-commerce growth partly reflects changes in the number of participating businesses rather "
      "than only changes in transaction value per business.")
w.para("This separates changes associated with wider business participation from changes in the average value "
       "observed per business.")
w.hyp("H3. Marketplace participation is associated with stronger financial recordkeeping.")
w.para("Association does not imply that marketplace participation causes formalization.")
w.hyp("H4. Measurement choices materially alter conclusions about the scale and growth of digital economic "
      "activity.")
w.caption("Table 1. Hypotheses and principal empirical tests.")
w.table([
    ["Hypothesis", "Principal empirical test"],
    ["H1 - Transaction-revenue divergence",
     "Matched platform transaction value and revenue over time, the Ecosystem Ratio, and annual growth divergence"],
    ["H2 - Sources of e-commerce growth", "BPS national e-commerce value and business-count decomposition"],
    ["H3 - Marketplace participation and recordkeeping", "BPS published marketplace and financial-recordkeeping evidence"],
    ["H4 - Measurement choice", "Comparison of platform, national e-commerce, and payment records"],
])
w.para("The separate comparison of platform businesses outside Indonesia is used as corroboration rather than as a "
       "fifth hypothesis. Indonesia's marketplace reporting regime likewise remains an institutional application. "
       "Whether platform-held records become administratively linked to seller obligations requires "
       "post-implementation evidence and cannot be established from the legal framework alone.")

w.h1("4. Data and Methodology")
w.para("The study uses several sources because no single record captures every part of digital platform activity. "
       "Platform reports provide transaction value and revenue; BPS-Statistics Indonesia provides national "
       "e-commerce and business evidence; Bank Indonesia provides payment statistics; and platform companies "
       "outside Indonesia provide a separate comparison. These sources answer different questions and are not "
       "combined into one dataset.")
w.h2("4.1 Selecting the Platform Evidence")
w.para("The basic observation is one platform in one year, with transaction value and revenue referring to the same "
       "period and, as closely as possible, the same business activity.")
w.para("An observation is retained only when its period, geographic and business coverage, units, definitions, and "
       "source can be identified. Measures referring to different periods are excluded, while major changes in "
       "accounting definitions or business structure are recorded.")
w.caption("Table 2. Indonesian platform evidence.")
w.table([
    ["Type of evidence", "Platforms", "Annual observations", "Role"],   # EDIT: "Platform series" -> "Platforms"
    ["Direct Indonesia-aligned data", "Tokopedia e-commerce", "2", "Main direct anchor"],
    ["Company-reported data with broader scope", "Blibli third-party business; Bukalapak Group", "9",
     "Evidence through time with stated scope limits"],
    ["Indonesia estimates requiring construction", "Grab; Shopee Indonesia", "6", "Supporting country evidence"],
])
w.para("Table 2 groups the retained evidence by how directly it maps to Indonesia. An annual observation means one "
       "matched transaction-and-revenue pair for one platform in one year, and the 17 retained observations produce "
       "12 year-to-year comparisons because the same platforms are not available in every year.")
w.para("Tokopedia is closely aligned with Indonesia but is not a literal country line. Blibli includes online travel, "
       "Bukalapak includes some overseas activity, and Grab and Shopee require additional construction.")
w.h2("4.2 Comparing Platform Activity Through Time")
w.para("Changes through time are examined mainly by comparing each platform with itself. This keeps the company, "
       "business activity, and reporting convention as similar as the disclosures allow.")
w.para("When transaction value and revenue move sharply apart, disclosed changes in monetization, customer "
       "incentives, revenue recognition, and business scope are used to reconcile the movement where the filings "
       "permit. These reconciliations are arithmetic accounting explanations rather than estimates of causal "
       "effects.")
gtv = w.para("Companies use different names for transaction activity, including gross transaction value (GTV), gross "
             "merchandise value (GMV), and total payment volume (TPV). This proposal uses GTV as its common label "
             "for the value of transactions processed through a platform, and each issuer's own label and definition "
             "are retained and checked rather than assumed to be identical.")
footnote_ref(gtv, 2)
w.para("A separate comparison uses the 2023 fiscal year (FY2023) because it provides the most useful common year "
       "across the principal platform cases. It shows the scale of the transaction-revenue difference among the "
       "selected cases, not the size of Indonesia's entire digital economy. Market-share estimates describe "
       "transaction-value coverage only, not wider representativeness.")
w.h2("4.3 Direct and Constructed Indonesia Measures")
w.para("Not every company reports the same Indonesia-specific information. Table 3 sets out, for each platform, what "
       "is disclosed directly, what has to be constructed, and the limitation that follows.")
w.caption("Table 3. Construction of the principal platform measures.")
w.table([
    ["Platform", "Directly available", "Constructed or external input", "Main limitation"],
    ["Tokopedia", "E-commerce transaction value and segment revenue", "None for matched pair",
     "Indonesia-aligned segment, not literal country line"],
    ["Grab", "Indonesia revenue; Group revenue and GTV", "Indonesia GTV derived from Group figures",
     "Assumes Group revenue-to-GTV relationship applies to Indonesia"],
    ["Shopee", "Group service monetization rate", "External Indonesia transaction estimate; Indonesia revenue derived",
     "Country measures are not independent disclosures"],
    ["Blibli / Bukalapak", "Reported transaction and revenue pairs", "None for retained pairs",
     "Scope extends beyond Indonesia alone"],
])
w.para("Grab reports Indonesian revenue directly but not Indonesian transaction value. Indonesian GTV is therefore "
       "estimated as:")
w.eq("Indonesia GTV = Indonesia revenue × Group GTV / Group revenue")
w.para("The Group monetization rate is platform revenue divided by gross transaction value, and the result is treated "
       "as a reconstruction rather than a direct disclosure.")
w.para("Sea Limited does not report Shopee transaction value separately for Indonesia. Indonesian transaction value "
       "is taken from Momentum Works, a Singapore-based industry research firm whose annual Southeast Asian "
       "e-commerce report is widely cited for country-level platform transaction estimates. It is used because no "
       "issuer discloses Shopee's Indonesian transaction value, it is labeled as a third-party estimate wherever it "
       "appears, and Indonesian revenue is then estimated using Sea's Group service monetization rate:")
w.eq("Indonesia revenue = Indonesia transaction value × Group monetization rate")
w.para("Both constructions remain explicitly labeled wherever they are used.")
w.h2("4.4 Evidence Beyond Platform Accounts")
w.caption("Table 4. Evidence modules and their role.")
w.table([
    ["Evidence source", "What is observed", "Role"],
    ["Indonesian platform histories", "Transaction value and revenue through time", "Measure the wedge and its movement"],
    ["BPS national e-commerce statistics", "National transaction value and estimated number of businesses",
     "Examine sources of aggregate growth"],
    ["BPS marketplace study", "Business-level marketplace and recordkeeping evidence",
     "Examine the association with financial recordkeeping"],
    ["Bank Indonesia", "Digital-payment measures", "Compare payment growth with e-commerce growth"],
    ["Platform companies outside Indonesia", "Matched transaction and revenue measures",
     "Test whether the same boundary appears elsewhere"],
])
w.para("Table 4 gives each source a single role. The modules answer different questions, so they are reported "
       "separately rather than pooled into one estimate.")
w.para("For H2, national e-commerce value is expressed as:")
w.eq("V = N × A")
w.para("where N is the estimated number of e-commerce businesses and A is implied nominal transaction value per "
       "business. Changes in total value are divided arithmetically between the business-count and "
       "value-per-business components. This does not identify causal entry, productivity, or growth among "
       "continuing firms. A conflicting published 2023 business count is retained as a sensitivity check.")
w.para("For H3, BPS's published business-level evidence is used to examine the association between marketplace "
       "participation and financial recordkeeping. Province-level correlations are supporting diagnostics only.")
w.para("Bank Indonesia payment measures provide a separate comparison for H4. They are compared with e-commerce "
       "growth but are not treated as sales estimates or added to platform or BPS transaction values.")
w.h2("4.5 External Evidence and Robustness")
w.para("The external comparison contains 48 matched annual transaction-and-revenue observations from eight platform "
       "businesses covering 2017-2025. It tests whether the same transaction-revenue boundary appears under other "
       "business models; it is not a representative sample of global e-commerce and does not enlarge the Indonesian "
       "sample.")
w.para("Robustness checks vary admission rules, revenue definitions, publication versions, and assumptions used in "
       "reconstructed country measures. The resulting ranges are treated as sensitivity or scenario ranges rather "
       "than statistical confidence intervals.")

w.h1("5. Preliminary Evidence and Feasibility")
w.para("The evidence assembled so far is sufficient to establish that the proposed comparisons can be carried out and "
       "that the choice of measurement produces economically meaningful differences. The preliminary results first "
       "show the scale and movement of the transaction-revenue boundary, then examine whether the wider Indonesian "
       "evidence tells the same story.")
w.h2("5.1 Platform Scale and Movement")
w.para("The 2023 fiscal year provides a common comparison across Grab, Tokopedia, and Shopee.")
w.caption("Table 5. Indonesia-focused 2023 platform comparison, US$ billions.")
w.table([
    ["Case", "Transaction value V", "Revenue R", "Wedge W", "Ecosystem Ratio E", "Evidence"],
    ["Grab Indonesia", "5.381", "0.605", "4.776", "7.895x", "Transaction value reconstructed"],
    ["Tokopedia", "16.331", "0.405", "15.926", "39.296x", "Direct matched pair"],
    ["Shopee Indonesia", "21.520", "2.152", "19.368", "9.000x", "Transaction value and revenue estimated"],
    ["Selected cases", "43.233", "3.162", "40.070", "12.671x", "Sum of three cases, not national total"],
])
w.para("Table 5 reports that cross-section. Across the three cases, US$43.23 billion of transaction value corresponds "
       "to US$3.16 billion of recognized platform revenue, leaving a US$40.07 billion wedge, so transaction value is "
       "approximately 13.7 times recognized revenue. For scale, that wedge is equivalent in magnitude to about 2.9 "
       "percent of Indonesia's 2023 nominal GDP of US$1.371 trillion. The comparison is one of magnitude only: the "
       "three platforms are not the whole market, and the wedge is transaction value outside platform revenue, not "
       "value added.")
# EDIT: the one assumption an examiner will test first, with its preliminary sensitivity
w.para("Shopee's ratio of 9.0 follows from the Group monetization rate used to estimate its Indonesian revenue. In "
       "preliminary checks, rates between 9 and 11 percent move the combined wedge only between US$39.86 billion "
       "and US$40.29 billion, and excluding any one platform leaves a wedge of at least US$20.70 billion.")
w.para("The longitudinal evidence shows that this relationship is not constant. Across twelve year-to-year "
       "comparisons, revenue grows faster than transaction value in ten, with a median absolute growth divergence "
       "of 42.02 percentage points. Figure 1 tracks the Ecosystem Ratio for each platform across those years; it "
       "ends lower than it started in every series, though not by a steady decline.")
w.element(figure_el)
w.caption("Figure 1. Ecosystem Ratio through time for each retained platform. Each line follows one platform's "
          "matched figures; differences between platforms are descriptive.", center=True)
w.para("Tokopedia provides the clearest illustration of why the two measures can separate. Between 2022 and 2023, "
       "transaction value falls 8.9 percent while third-party net revenue rises 53.2 percent. The change in net "
       "revenue can be reconciled arithmetically to lower customer incentives and higher gross revenue rather than "
       "treated as an unexplained residual.")
w.h2("5.2 Evidence Beyond Platform Revenue")
w.para("National e-commerce and business evidence provides a different view of the same digital economy. Table 6 "
       "summarizes the preliminary evidence corresponding to H2-H4.")
w.caption("Table 6. Preliminary evidence beyond platform accounts.")
w.table([
    ["Question", "Preliminary evidence", "Interpretation"],
    ["Where does e-commerce growth come from?",
     "E-commerce value +17.08%; estimated business count +15.31%; implied nominal value per business +1.54%",
     "Much of the aggregate increase is associated with broader participation rather than a large increase in "
     "value per business"],
    ["Where does that growth occur?",
     "Approximately 98.46% of the nominal increase occurs outside the marketplace component",
     "Marketplace-centered evidence captures only part of recent national e-commerce growth"],
    ["Are marketplace use and financial records related?",
     "Complete financial statements: 28.63% of marketplace users versus 12.25% of non-marketplace users",
     "Marketplace participation is associated with stronger financial recordkeeping; the evidence is observational"],
    ["Do other digital records show the same growth?",
     "BPS e-commerce +17.08%; electronic-money shopping +30.47%; mobile-banking payments and purchases +82.84%; "
     "QRIS +186.98%",
     "Different records produce substantially different descriptions of digital growth"],
])
w.para("These comparisons are deliberately kept separate. The business-count decomposition is arithmetic rather than "
       "causal, marketplace participation is treated as an association rather than a cause of formalization, and "
       "payment values are not interpreted as alternative estimates of e-commerce sales.")
w.para("The evidence nevertheless shows why measurement choice matters. Recent e-commerce growth is concentrated "
       "differently depending on whether the analyst looks at marketplace activity, national e-commerce statistics, "
       "or payment records.")
w.h2("5.3 Feasibility and Scope of the Preliminary Findings")
w.para("The preliminary evidence is sufficient to motivate the full analysis. H1 is supported by substantial movement "
       "between transaction value and revenue within platforms. H2 and H3 can be examined using national and "
       "business-level BPS evidence, while the contrast between platform, e-commerce, and payment records provides a "
       "direct test of H4.")
# EDIT: implementation date and the regime's coverage, a design fact rather than an outcome
w.para("Indonesia's marketplace reporting regime provides an institutional application of the same distinction. PMK "
       "37/2025, administered by the Directorate General of Taxes (DJP), shows how seller-linked transaction records "
       "can enter reporting and withholding, but implementation "
       "begins only on 1 November 2026, so post-implementation outcomes are not yet available. Because the regime "
       "designates marketplace operators, its reach is bounded by the marketplace channel, which carried 15.8 percent "
       "of 2024 e-commerce value and 1.5 percent of 2023-2024 growth. The regulation is therefore examined as "
       "institutional architecture, not as evidence of unpaid tax or of a completed compliance effect.")
w.para("At this stage, the preliminary results establish that the proposed measurement can be constructed, followed "
       "through time, connected to disclosed accounting changes, and compared with independent records of Indonesian "
       "digital activity.")

w.h1("6. Why the Problem Matters")
w.para("These records are designed to measure different things, and the problem arises when one is used as a "
       "substitute for another. Platform revenue measures what the platform recognizes "
       "as its own income, transaction value measures the commerce it processes, national e-commerce statistics "
       "describe a broader market, and payment systems record the movement of funds. When these measures move "
       "differently, the choice of record can change conclusions about economic scale, growth, and monetization.")
w.para("Administrative use is a separate question. Platform-held transaction records may contain information about "
       "sellers and transactions that does not appear in the platform's corporate revenue, but those records become "
       "useful to government only when identity, reporting, transmission, and matching are in place. PMK 37/2025 "
       "illustrates this institutional link without establishing a compliance or tax-revenue effect.")
w.para("The thesis therefore treats measurement and administrative visibility as related but distinct problems.")

w.h1("7. Limitations and Robustness")
w.para("The evidence supports several conclusions, but the interpretation of each result is deliberately bounded. "
       "Table 7 pairs each supported conclusion with what it does not establish.")
w.caption("Table 7. Main inference boundaries.")
w.table([
    ["Evidence supports", "Does not by itself establish"],
    ["W, E, and D describe the size and movement of the transaction-revenue boundary.",
     "Participant profit, taxable income, unpaid tax, tax evasion, or missing GDP."],
    ["Separate evidence types show how directly each platform measure represents Indonesia.",   # EDIT: label
     "That direct, broader-scope, and reconstructed observations are interchangeable."],
    ["Disclosed revenue components can reconcile selected movements in the wedge.",
     "A causal effect of incentives, monetization, or accounting choices."],
    ["BPS evidence describes aggregate e-commerce growth and recordkeeping associations.",
     "Causal business entry, productivity effects, or marketplace-induced formalization."],
    ["Payment statistics and PMK provide additional measurement and institutional evidence.",
     "Equivalent measures of sales or a completed compliance or tax-revenue effect."],
])
w.para("Robustness analysis tests whether the main findings depend on one preferred construction. The longitudinal "
       "results are compared under narrower evidence-admission rules, alternative revenue definitions, different "
       "publication versions, and alternative assumptions for reconstructed country measures. These checks are "
       "reported as sensitivity analyses rather than statistical confidence intervals.")
w.h1("8. Next Steps")
w.para("The main data and empirical framework are already assembled. The remaining work is to complete the robustness "
       "checks, finalize the treatment of the different evidence types, verify source consistency, and develop the "
       "preliminary results into the full thesis analysis. Table 8 sets out that schedule.")
w.caption("Table 8. Planned work toward the final thesis.")
w.table([
    ["Period", "Planned work"],
    ["Proposal stage", "Incorporate feedback from the proposal examination and finalize the empirical specification."],
    ["October-November 2026", "Complete remaining source checks and robustness analysis; finalize the main tables and figures."],
    ["Before the final defense", "Complete the full manuscript, integrate supporting evidence where it strengthens the "
     "argument, and conduct the final citation and consistency review."],
])
w.para("The proposed study begins from a simple distinction between the value a platform processes and the revenue it "
       "recognizes. The preliminary evidence indicates that this boundary is economically large, changes through "
       "time, and produces different conclusions when compared with other records of digital activity. The final "
       "thesis will test how robust those conclusions remain under alternative evidence and measurement choices.")

w.h1("Appendix A - Data Sources by Evidence Type")
w.caption("Table A1. Source lineage by evidence type.")
w.table([
    ["Evidence type", "Platform / source", "Transaction-value source", "Revenue / comparison source"],
    ["Direct Indonesia-aligned", "Tokopedia e-commerce", "GoTo annual report, e-commerce segment metrics",
     "GoTo annual report, segment note"],
    ["Company-reported broader scope", "Blibli third-party business", "Global Digital Niaga prospectus and results releases",
     "Same issuer filings"],
    ["Company-reported broader scope", "Bukalapak Group", "Bukalapak annual and sustainability reports", "Same issuer filings"],
    ["Indonesia reconstruction", "Grab Indonesia", "Derived from Indonesia revenue and Group monetization rate",
     "Grab Form 20-F"],
    ["Indonesia reconstruction", "Shopee Indonesia", "Momentum Works Southeast Asia e-commerce estimate",
     "Derived from Sea Limited Form 20-F"],
    ["National statistics", "BPS-Statistics Indonesia", "E-Commerce Statistics 2023 and 2024; marketplace study",
     "Business-count and recordkeeping evidence"],
    ["Payment statistics", "Bank Indonesia", "Payment System Statistics and QRIS reports", "Growth-rate comparison only"],
    ["External corroboration", "Eight platform businesses", "Issuer transaction measures", "Issuer revenue measures"],
    ["Regulatory", "PMK 37/2025; DJP", "Ministry of Finance / Directorate General of Taxes", "Institutional architecture only"],
])
w.para("Table A1 traces each evidence type to its sources. Every figure used in this proposal traces to a named file "
       "and locator in the accompanying data package.")

add_footnotes(doc, ["Issuer labels differ. GoTo reports GTV and Sea Limited reports GMV for Shopee; for these the "
                    "label differs but the measure is the same. Blibli and Bukalapak report TPV for the segments "
                    "used here, and Grab reports GMV for deliveries and mobility and TPV for financial services. "
                    "Where a source reports TPV the measure rests on a different basis, so it is retained under its "
                    "own label rather than relabeled as GTV."])
doc.save(OUT)
print("wrote", OUT)
