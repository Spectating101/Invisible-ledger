# -*- coding: utf-8 -*-
import sys, re
sys.path.insert(0,'/tmp/build')
from surg_lib import *
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
B='/home/phyrexian/Downloads/llm_automation/project_portfolio/Solarpunk-bitcoin/'
F=B+'Invisible_Ledger_Proposal_Sept01_Surgery_2026-09-14.docx'
d=Document(F); body=d.element.body
def acc(ch):
    x=re.sub(r'<w:del [^>]*>.*?</w:del>','',ch.xml,flags=re.S)
    return ' '.join(re.sub(r'<[^>]+>','',x).split())
def find(pred):
    for ch in body.iterchildren():
        if ch.tag.endswith('}p') and pred(acc(ch)): return ch
    return None
def ins_after(a, style, text): return new_paragraph(body, style, [('ins',text)], a)
def ins_table(anchor, header, rows, size='19'):
    t=d.add_table(rows=1, cols=len(header))
    try: t.style='Table Grid'
    except Exception: pass
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    trPr=t.rows[0]._tr.get_or_add_trPr()
    th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    def cellset(cell, txt, bold):
        cell.text=''; p=cell.paragraphs[0]._p
        r=_run(txt); rpr=OxmlElement('w:rPr')
        rf=OxmlElement('w:rFonts')
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
        rpr.append(rf); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),size); rpr.append(sz)
        if bold: rpr.append(OxmlElement('w:b'))
        r.insert(0,rpr); p.append(_wrap('ins',[r]))
    for i,h in enumerate(header): cellset(t.rows[0].cells[i],h,True)
    for row in rows:
        tr=t.add_row(); trp=tr._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        ins=OxmlElement('w:ins'); ins.set(qn('w:id'),_id()); ins.set(qn('w:author'),AUTHOR); ins.set(qn('w:date'),DATE)
        trp.append(ins)
        for i,v in enumerate(row): cellset(tr.cells[i],v,False)
    body.remove(t._tbl); anchor.addnext(t._tbl)
    return t._tbl

a=find(lambda t:t.startswith('BPS-Statistics Indonesia evidence, reported in Section 5.5'))

a=ins_after(a,'H1','5. Preliminary Evidence and Feasibility')
a=ins_after(a,'BodyFirst','The retained evidence already supports each step of the design: the wedge can be measured, it moves, its movement can be traced to disclosed revenue components, the pattern recurs outside Indonesia, and official statistics locate activity beyond platform accounts. These are feasibility diagnostics for the proposal examination, not final estimates or causal findings, and the final sample admission remains subject to advisor approval.')

# 5.1 FY2023 cross-section
a=ins_after(a,'H2','5.1 FY2023 cross-section')
a=ins_after(a,'BodyFirst','FY2023 is retained as the cleanest common cross-platform snapshot, not as the period of analysis. It is the year for which a detailed Indonesian marketplace-share denominator is available for all three leading platforms. Table 6 reports the wedge and the Ecosystem Ratio for each case and for the three combined.')
cap=ins_after(a,'TableCaption','Table 6. Indonesia-focused FY2023 cross-section. Values in US$ billions; V is transaction value, R is platform revenue, E is the Ecosystem Ratio. Evidence class: Grab\u2019s transaction value is derived from its Group monetization rate; Tokopedia is a directly reported Indonesia-aligned segment pair; both Shopee inputs combine an external market estimate with a disclosed company-wide rate. The final row sums three documented cases under mixed evidence classes and is not an Indonesia-wide total.')
T=ins_table(cap,["Case","V","R","Wedge","E","Evidence class"],
 [["Grab Indonesia","5.381","0.605","4.776","7.895×","Derived V"],
  ["Tokopedia e-commerce","16.331","0.405","15.926","39.296×","Direct pair"],
  ["Shopee Indonesia","21.520","2.152","19.368","9.000×","Derived V and R"],
  ["Selected platforms","43.233","3.162","40.070","12.671×","Sum, not a total"]])
a=new_paragraph(body,'Body',[('ins','The selected-platform wedge of US$40.07 billion is equivalent to approximately 2.9 percent of Indonesia’s 2023 nominal gross domestic product. Because two of the three cases depend on an assumption, every assumed parameter is varied one at a time and the result is also reported with each platform excluded in turn. Varying Grab’s implied Indonesia monetization rate moves the selected-platform wedge between US$39.17 billion and US$41.42 billion; varying Shopee’s assumed Indonesian market share from 35 to 45 percent gives US$37.65 billion to US$42.49 billion; varying Shopee’s revenue-to-GMV rate gives US$39.86 billion to US$40.29 billion; and varying the Indonesian e-commerce market denominator gives US$38.13 billion to US$42.01 billion. Excluding Grab leaves a wedge of US$35.29 billion across the two remaining cases, excluding Tokopedia leaves US$24.14 billion, and excluding Shopee leaves US$20.70 billion. No single assumption or platform drives the order of magnitude. Reproduction checks on the sample, the applied rate and the aggregate are retained in the data package. These checks do not convert three documented cases into a statistical sample.')],T)

# 5.2 two universes
a=ins_after(a,'H2','5.2 Longitudinal divergence under two evidence universes')
a=ins_after(a,'BodyFirst','Two valid longitudinal summaries coexist because they admit different evidence. Both are reported. Table 7 reports both. The all-tier inventory includes the conditional Grab and Shopee reconstructions; the direct-candidate sensitivity excludes them and instead extends Blibli with its directly transcribed FY2020 prospectus observation.')
cap=ins_after(a,'TableCaption','Table 7. Within-series annual growth divergence under two admission rules. The counts differ because evidence admission differs; the difference is reported rather than resolved into a single headline.')
T=ins_table(cap,["","All-tier inventory","Direct-candidate sensitivity"],
 [["Series admitted","Blibli FY2021–25, Bukalapak, Tokopedia, Grab, Shopee","Blibli FY2020–25, Bukalapak, Tokopedia"],
  ["Transitions","12","9"],
  ["Revenue grows faster","10","6"],
  ["Transaction grows faster","2","3"],
  ["Sign reversals","2","3"],
  ["Median absolute divergence","42.02 pp","42.94 pp"]])
a=new_paragraph(body,'Body',[('ins','The additional sign reversal in the direct-candidate universe is Blibli FY2020→FY2021. That the reversal count changes with admission is itself informative: it demonstrates on the study’s own data the proposition that measurement choices can alter an empirical conclusion, which is why the sample hierarchy is shown rather than collapsed.')],T)

# 5.3 mechanism
a=ins_after(a,'H2','5.3 Mechanism: what moves the wedge')
a=ins_after(a,'BodyFirst','Tokopedia FY2022–FY2023 is the strongest direct case. Transaction value falls 8.90 percent while third-party net segment revenue rises 53.20 percent. Reconciling the reported revenue components, approximately 60.56 percent of the increase in net revenue is associated with lower customer incentives and 39.44 percent with higher gross revenue; the disclosed components sum exactly to reported net revenue in both years. Recognized revenue therefore strengthened while facilitated activity weakened, for reasons visible in the issuer’s own disclosure. A second, weaker case points the same way. Blibli 3P Retail FY2022–FY2023 shows transaction value rising 34.72 percent while segment net revenue rises 465.33 percent. That divergence is arithmetically real but arises from a very low base — segment net revenue moves from IDR 199 billion to IDR 1,125 billion, about US$13 million to US$74 million — and the segment perimeter also contains travel, so it is reported as evidence of a monetization-regime change rather than as a measurement of the same quantity as the Tokopedia case.')

# 5.4 global
a=ins_after(a,'H2','5.4 Is the wedge specific to Indonesia?')
a=ins_after(a,'BodyFirst','A separate module tests whether the same measurement boundary appears in platform businesses outside Indonesia. Across 48 matched issuer-years for eight businesses — eBay, Etsy, Shopify, Jumia, Zalando, Rakuten, Mercado Libre, and Sea — the transaction-to-revenue relationship is present throughout, but its level is business-model determined: take rates range from 0.22 percent to 74.63 percent. Of 40 annual transitions, 29 are clean-scope, with a median absolute growth divergence of 7.10 percentage points and four transitions in which transaction value and revenue move in opposite directions. The wedge is therefore a general property of platform intermediation rather than an Indonesian artefact, while its magnitude is not a stable parameter across business models. Issuers without a continuous matched transaction series — Amazon, Alibaba, PDD, and Coupang — are excluded and the reason recorded.')

# 5.5 BPS + institutional
a=ins_after(a,'H2','5.5 Official statistics and institutional evidence')
a=ins_after(a,'BodyFirst','BPS evidence locates digital commercial activity beyond platform accounts. National e-commerce value rises from Rp1,100.87 trillion in 2023 to Rp1,288.93 trillion in 2024, approximately US$72.2 billion to US$81.4 billion, an increase of 17.08 percent, while estimated e-commerce businesses rise from 3,816,750 to 4,400,972, an increase of 15.31 percent; implied nominal value per estimated business therefore rises only 1.54 percent. The channel split is more striking: marketplace value rises 1.45 percent against 20.57 percent for the non-marketplace component, so approximately 98.46 percent of the nominal increase falls outside the marketplace component that platform accounts observe. National financial-report ownership among e-commerce businesses is 17.15 percent in 2024. BPS also publishes an analysis reporting higher financial-report ownership among marketplace users than non-users; that is BPS’s own published result and is cited as such. Province-level evidence is ecological and is not used to infer a business-level relationship.')
a=ins_after(a,'Body','Taken together, the modules show that the wedge can be measured under explicit admission rules, that it moves for identifiable reasons, that the pattern is not specific to Indonesia, that most recent national growth occurs outside the channel platform accounts observe, and that a legal route now exists from marketplace records into tax administration. The research question can be answered with evidence that is available, tiered, and auditable to source.')


# ---- 6. Limitations ----
a=ins_after(a,'H1','6. Limitations and Advisor Decisions')
a=ins_after(a,'BodyFirst','The principal limitation is that no Indonesian platform publishes a strictly country-labelled matched pair of transaction value and revenue. Every observation is therefore admitted under a named evidence tier, and the tiers differ in what they can support. Tokopedia is an Indonesia-aligned business segment rather than a country line. Blibli’s 3P Retail perimeter includes online travel. Bukalapak’s Group figures include overseas operations. Grab’s Indonesian transaction value is derived from a Group monetization rate, and both Shopee inputs combine an external market estimate with a disclosed company-wide rate. These are properties of what issuers disclose, not choices made for convenience, and they bound what the study can claim.')
a=ins_after(a,'Body','Three further boundaries apply. Mechanism reconciliations are arithmetic against the issuer’s own disclosed categories, not causal decompositions. BPS province evidence is ecological and cannot identify a business-level relationship. PMK 37/2025, postponed to 1 November 2026, offers no completed period from which to estimate a compliance or revenue effect.')
a=ins_after(a,'Body','Two decisions are sought at the proposal examination. First, whether the direct issuer scope-pending tier — Blibli 3P Retail and Bukalapak Group — should be admitted to the final longitudinal main sample, reported only as a labelled tier, or excluded. Second, whether the conditional country reconstructions for Grab and Shopee may appear in the main comparison or should be confined to sensitivity analysis. Both decisions change the reported transition counts, as Table 7 shows, and neither can be settled by the data alone.')

# ---- 7. Work plan ----
a=ins_after(a,'H1','7. Work Plan')
a=ins_after(a,'BodyFirst','Table 8 sets out the remaining work. The empirical backend is assembled and source-linked, so the outstanding tasks are sample admission, any analysis the committee requests, and construction of the final manuscript under the approved boundary.')
cap=ins_after(a,'TableCaption','Table 8. Planned work to the final defence. The empirical backend is assembled and source-linked; the remaining work is admission, requested analysis, and manuscript construction under the approved boundary.')
ins_table(cap,["Period","Planned work"],
 [["September 2026","Advisor review of the data package; proposal oral examination; freeze the final sample-admission rule and the checks the committee requests."],
  ["October–November 2026","Complete requested mechanism reconciliations and source concordance; extend the sensitivity analysis to the approved tier boundary; incorporate committee feedback; freeze thesis tables and figures."],
  ["Before final defence","Rebuild the full manuscript under the approved sample; retain the global, official-statistics and institutional modules only where they advance the final argument."]])


# ---- Appendix A: variable definitions and data sources (advisor comment 153) ----
a=ins_after(a,'H1','Appendix A. Definition of Variables and Data Sources')
cap=ins_after(a,'TableCaption','Table A1. Variables used in the analysis, their definitions, and the source of each. Transaction value and platform revenue are observed from issuer disclosures; the wedge, the Ecosystem Ratio, the growth rates and their difference are derived from unrounded source inputs.')
ins_table(cap,["Variable","Definition","Source"],
 [["V","Gross transaction value processed through the platform in the period, at the stated business and geographic scope","Issuer annual reports, prospectuses and results materials"],
  ["R","Platform revenue recognized for the same period and scope","Issuer annual reports and financial statements"],
  ["W = V − R","Absolute invisible wedge: transaction value not booked as platform revenue","Derived"],
  ["E = (V − R) / R","Ecosystem Ratio: the wedge relative to booked platform revenue","Derived"],
  ["g(V), g(R)","Year-on-year growth in transaction value and in platform revenue","Derived"],
  ["D = g(V) − g(R)","Growth divergence between transaction value and platform revenue","Derived"],
  ["Take rate","Platform revenue divided by gross transaction value","Derived"],
  ["Revenue components","Customer incentives, third-party gross revenue, third-party net revenue","Issuer segment notes where separately disclosed"],
  ["Evidence tier","Direct Indonesia-aligned, direct issuer scope-pending, or conditional country reconstruction","Assigned by the admission rules in Section 4.2"],
  ["E-commerce value","National nominal e-commerce transaction value","BPS-Statistics Indonesia, E-Commerce Statistics"],
  ["E-commerce businesses","Estimated number of businesses conducting e-commerce","BPS-Statistics Indonesia, E-Commerce Statistics"],
  ["Financial-report ownership","Share of e-commerce businesses holding financial statements","BPS-Statistics Indonesia, national indicators"]])

d.save(F); print('stage3b: §5 + §6 + §7 + Appendix A built')
