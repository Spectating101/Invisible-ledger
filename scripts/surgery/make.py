# -*- coding: utf-8 -*-
import sys, copy
sys.path.insert(0,'/tmp/build')
from content import *
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC="/home/phyrexian/Downloads/llm_automation/project_portfolio/Solarpunk-bitcoin/Invisible_Ledger_Proposal_FLASHPOINT_GRAFT_PLAIN_FINAL_2026-09-13.docx"
OUT="/home/phyrexian/Downloads/llm_automation/project_portfolio/Solarpunk-bitcoin/Invisible_Ledger_Proposal_Kong_Rootstock_2026-09-14.docx"

d=Document(SRC)
body=d.element.body

# --- update cover title (paragraphs 5 = ZH, 6 = EN) ---
def set_text(p, txt):
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if p.runs: p.runs[0].text = txt
paras=d.paragraphs
set_text(paras[5], TITLE_ZH)
set_text(paras[6], TITLE_EN)

# --- delete everything after the cover page break (index 13) ---
kids=list(body.iterchildren())
for ch in kids[14:]:
    if ch.tag.endswith('}sectPr'): continue
    body.remove(ch)

sectPr = body.find(qn('w:sectPr'))

def add(style=None):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    if sectPr is not None:
        body.remove(p._element); sectPr.addprevious(p._element)
    return p

def enforce(p, size=12, bold=None, italic=None, align=None, space_after=3, space_before=0, keep=False):
    pf=p.paragraph_format
    pf.space_after=Pt(space_after); pf.space_before=Pt(space_before)
    pf.line_spacing=1.0583333333333333
    if align is not None: p.alignment=align
    if keep: pf.keep_with_next=True
    for r in p.runs:
        r.font.name='Times New Roman'; r.font.size=Pt(size)
        r.font.color.rgb=RGBColor(0,0,0)
        if bold is not None: r.font.bold=bold
        if italic is not None: r.font.italic=italic
        rpr=r._element.get_or_add_rPr()
        rf=rpr.find(qn('w:rFonts'))
        if rf is None:
            rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
            rf.set(qn(a),'Times New Roman')

def para(text, style=None, size=12, bold=None, italic=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=3, space_before=0, keep=False):
    p=add(style); p.add_run(text)
    enforce(p,size,bold,italic,align,space_after,space_before,keep)
    return p

def shade(cell,hexv):
    tcPr=cell._tc.get_or_add_tcPr()
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexv)
    tcPr.append(sh)

def table(spec):
    hdr=spec['header']; rows=spec['rows']
    t=d.add_table(rows=1, cols=len(hdr)); t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=True
    # repeat header row across pages
    trPr=t.rows[0]._tr.get_or_add_trPr()
    th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i]; c.text=''
        p=c.paragraphs[0]; p.add_run(h)
        enforce(p,9.5,bold=True,align=WD_ALIGN_PARAGRAPH.LEFT,space_after=1,space_before=1)
        shade(c,'FFFFFF')
    for row in rows:
        tr=t.add_row(); cp=tr._tr.get_or_add_trPr()
        cs=OxmlElement('w:cantSplit'); cp.append(cs)
        cells=tr.cells
        for i,v in enumerate(row):
            cells[i].text=''
            p=cells[i].paragraphs[0]; p.add_run(v)
            enforce(p,9.5,bold=False,align=WD_ALIGN_PARAGRAPH.LEFT,space_after=1,space_before=1)
    if sectPr is not None:
        body.remove(t._tbl); sectPr.addprevious(t._tbl)
    return t

# ---- build body ----
for kind,val in BODY:
    if kind=='H1':
        para(val, style='Heading 1', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             space_before=8, space_after=4, keep=True)
    elif kind=='H2':
        para(val, style='Heading 2', size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             space_before=6, space_after=2, keep=True)
    elif kind=='P':
        para(val)
    elif kind=='EQ':
        para(val, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, space_before=4, space_after=6)
    elif kind=='TABLE':
        para(val['caption'], size=10, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             space_before=6, space_after=2, keep=True)
        table(val)
        pass

for r in REFS:
    p=para(r, space_after=2)
    p.paragraph_format.left_indent=Pt(18); p.paragraph_format.first_line_indent=Pt(-18)

para(PRIMARY_HEAD, style='Heading 1', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
     space_before=8, space_after=4, keep=True)
para(PRIMARY_NOTE)
for s in PRIMARY:
    p=para(s, space_after=2)
    p.paragraph_format.left_indent=Pt(18); p.paragraph_format.first_line_indent=Pt(-18)

# global font sweep incl. headers/footers
def sweep(container):
    for p in container.paragraphs:
        for r in p.runs:
            r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0,0,0)
            rpr=r._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
            if rf is None:
                rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
            for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
                rf.set(qn(a),'Times New Roman')
for s in d.sections:
    sweep(s.header); sweep(s.footer)

d.save(OUT)
print("saved", OUT)
