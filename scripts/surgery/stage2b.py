import sys, re
sys.path.insert(0,'/tmp/build')
from surg_lib import *
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
B='/home/phyrexian/Downloads/llm_automation/project_portfolio/Solarpunk-bitcoin/'
F=B+'Invisible_Ledger_Proposal_Sept01_Surgery_2026-09-14.docx'
d=Document(F); body=d.element.body
def acc(ch):
    x=re.sub(r'<w:del [^>]*>.*?</w:del>','',ch.xml,flags=re.S)
    return ' '.join(re.sub(r'<[^>]+>','',x).split())
def find(pred):
    for ch in body.iterchildren():
        if ch.tag.endswith('}p') and pred(acc(ch)): return ch
    return None
def ins_after(a,st,t): return new_paragraph(body,st,[('ins',t)],a)
def ins_table(anchor, header, rows, size='19'):
    t=d.add_table(rows=1, cols=len(header))
    try: t.style='Table Grid'
    except Exception: pass
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    trPr=t.rows[0]._tr.get_or_add_trPr()
    th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    def cs(cell,txt,bold):
        cell.text=''; p=cell.paragraphs[0]._p
        r=_run(txt); rpr=OxmlElement('w:rPr')
        rf=OxmlElement('w:rFonts')
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
        rpr.append(rf); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),size); rpr.append(sz)
        if bold: rpr.append(OxmlElement('w:b'))
        r.insert(0,rpr); p.append(_wrap('ins',[r]))
    for i,h in enumerate(header): cs(t.rows[0].cells[i],h,True)
    for row in rows:
        tr=t.add_row(); trp=tr._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        ins=OxmlElement('w:ins'); ins.set(qn('w:id'),_id()); ins.set(qn('w:author'),AUTHOR); ins.set(qn('w:date'),DATE)
        trp.append(ins)
        for i,v in enumerate(row): cs(tr.cells[i],v,False)
    body.remove(t._tbl); anchor.addnext(t._tbl)
    return t._tbl

a=find(lambda t:t.startswith('The thesis therefore asks how large the invisible wedge'))
a=ins_after(a,'Body','That motivating question is examined through the four linked research questions set out in Table 1. The first two establish the measurement and its movement; the third tests whether the pattern is specific to Indonesia; the fourth follows the activity and the records beyond platform accounts.')
a=ins_after(a,'TableCaption','Table 1. Research questions and principal evidence. RQ1 and RQ2 concern the measurement and its movement within Indonesia; RQ3 tests generality outside Indonesia; RQ4 follows the activity and the underlying records beyond platform accounts.')
ins_table(a,["","Research question","Principal evidence"],
 [["RQ1","How large is the invisible wedge in Indonesia, and does it persist across platforms and years?","17 admitted platform-year levels across three evidence tiers, FY2019–FY2025"],
  ["RQ2","What explains movements in the wedge?","Disclosed revenue components: incentives, gross and net revenue, monetization, business scope"],
  ["RQ3","Is the wedge specific to Indonesia or a general property of platform business models?","48 matched issuer-years across eight non-Indonesian platform businesses"],
  ["RQ4","Where does activity outside platform revenue appear, and can the underlying records reach administration?","BPS national, channel and recordkeeping evidence; PMK 37/2025 and DJP materials"]])

# ---- contributions 2-4 (the deleted paragraphs were never replaced) ----
a=find(lambda x:x.startswith('This study makes four contributions'))
a=ins_after(a,'Body','Second, it explains why the wedge moves. Transaction value and platform revenue can separate because of monetization decisions, customer incentives, gross-versus-net revenue recognition, or changes in reporting perimeter. The thesis reconciles the largest divergences against the revenue components issuers themselves disclose, so that movement in the wedge can be attributed either to platform economics or to accounting presentation rather than treated as a fixed structural ratio.')
a=ins_after(a,'Body','Third, it tests whether the wedge is specific to Indonesia. A separate module of forty-eight matched issuer-years across eight non-Indonesian platform businesses establishes whether the transaction-revenue boundary is a general property of platform intermediation and how far its magnitude varies with business model. This converts the wide cross-platform spread in measured ratios from an obstacle to comparison into a result in its own right.')
ins_after(a,'Body','Fourth, it connects the platform accounting boundary to participant records and tax administration in Indonesia. Statistics Indonesia evidence identifies where e-commerce activity and business recordkeeping occur beyond issuer accounts, while PMK 37/2025 and Directorate General of Taxes materials show how marketplace-held seller identity and transaction-linked turnover enter third-party reporting and withholding.')

d.save(F); print('1.2 RQs + contributions 2-4 inserted')
