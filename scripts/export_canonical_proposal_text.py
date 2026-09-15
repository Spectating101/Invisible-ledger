#!/usr/bin/env python3
"""Export a searchable Markdown mirror from the canonical proposal DOCX.

The FINAL DOCX is the proposal source of truth. This script exists only so code-search
and agents can retrieve current proposal text without consulting stale Markdown drafts.
Do not edit the generated Markdown by hand and do not rebuild the proposal from it.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "papers/current/Invisible_Ledger_Thesis_Proposal_FINAL_2026-09-14.docx"
OUTPUT = ROOT / "papers/current/Invisible_Ledger_Thesis_Proposal_CANONICAL_TEXT.md"


def iter_blocks(parent: _Document):
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def clean(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def esc_cell(text: str) -> str:
    return clean(text).replace("|", "\\|")


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"canonical proposal not found: {SOURCE}")

    source_bytes = SOURCE.read_bytes()
    source_sha = hashlib.sha256(source_bytes).hexdigest()
    doc = Document(SOURCE)

    lines: list[str] = [
        "<!-- GENERATED FILE: DO NOT EDIT. SOURCE OF TRUTH IS THE FINAL DOCX. -->",
        "# Invisible Ledger proposal — canonical searchable text",
        "",
        f"Generated from `{SOURCE.relative_to(ROOT)}`.",
        f"Source DOCX SHA-256: `{source_sha}`.",
        "",
        "> **Authority rule:** this file mirrors the current FINAL DOCX for search/review only. "
        "If this file and the DOCX ever disagree, the DOCX wins and this mirror must be regenerated.",
        "",
    ]

    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = clean(block.text)
            if not text:
                continue
            style = block.style.name if block.style is not None else ""
            if style == "Heading 1":
                lines.extend([f"## {text}", ""])
            elif style == "Heading 2":
                lines.extend([f"### {text}", ""])
            else:
                lines.extend([text, ""])
        else:
            rows = [[esc_cell(c.text) for c in row.cells] for row in block.rows]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * width) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")

    rendered = "\n".join(lines).rstrip() + "\n"
    OUTPUT.write_text(rendered, encoding="utf-8")
    print(f"wrote {OUTPUT.relative_to(ROOT)}")
    print(f"source sha256 {source_sha}")
    print(f"paragraphs {len(doc.paragraphs)} tables {len(doc.tables)}")


if __name__ == "__main__":
    main()
