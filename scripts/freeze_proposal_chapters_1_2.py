#!/usr/bin/env python3
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "papers/current/Invisible_Ledger_Thesis_Proposal_FINAL_2026-09-14.docx"
doc = Document(PATH)


def find_para(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise RuntimeError(f"paragraph not found: {text}")


def clear_between(start_text, end_text):
    start = find_para(start_text)._p
    end = find_para(end_text)._p
    node = start.getnext()
    while node is not None and node is not end:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def insert_before(anchor, text, style="Normal", *, bold=False, italic=False, center=False):
    el = OxmlElement("w:p")
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    p.style = style
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


# Chapter 1: approved working freeze. Abstract is deliberately untouched.
clear_between("1. Introduction", "2. Literature Review and Research Gap")
a2 = find_para("2. Literature Review and Research Gap")
ch1 = [
    ("Southeast Asia's digital economy grew from approximately US$100 billion in 2020 to US$263 billion in 2024 (Google, Temasek and Bain 2020, 2024). Digital platform scale, however, is not the same thing as platform revenue. Platforms record transactions, determine participant payouts, and book their own revenue, while national statistics and payment systems record other parts of the same economic activity. These measures are often used interchangeably even though they do not measure the same thing. This paper argues that doing so can produce systematically wrong conclusions about the scale and growth of platform economies.", "Normal", False, False, False),
    ("Suppose a platform processes 100 units of transaction value but recognizes only 10 as revenue. The remaining 90 are still recorded within the platform's system but do not constitute platform revenue; they may represent merchant receipts, driver payouts, inventory costs, taxes, and other pass-through payments. An analyst using the platform's accounts sees 10, while one measuring the commerce it processes sees 100. I call the difference the invisible wedge.", "Normal", False, False, False),
    ("Indonesia provides a useful empirical setting because several records of the same activity can be compared. Platform companies disclose transaction and revenue information; BPS-Statistics Indonesia, the national statistical agency, reports e-commerce activity; and Bank Indonesia, the country's central bank, reports digital payments. In 2023, the platform comparison covers Shopee and Tokopedia, which together represented about 70 percent of estimated Indonesian marketplace transaction value, together with Grab, a major intermediary in delivery and other digital services. The three cases processed US$43.23 billion of transaction value against US$3.16 billion of recognized revenue, leaving a US$40.07 billion invisible wedge.", "Normal", False, False, False),
    ("The difference also changes over time. Across twelve year-to-year comparisons between 2020 and 2025, revenue grew faster than transaction value in ten, with a median divergence of approximately 42 percentage points; in some years the two measures moved in opposite directions. Platform revenue therefore understates the level of platform-mediated commerce while, in most observed years, overstating its growth.", "Normal", False, False, False),
    ("The same problem appears beyond company accounts. BPS reports that national e-commerce value grew 17.1 percent between 2023 and 2024, but about 98.5 percent of that increase occurred outside the marketplace component; digital payment measures reported by Bank Indonesia grew substantially faster still. Indonesia's Minister of Finance Regulation No. 37 of 2025, which introduces seller reporting and withholding through designated marketplaces, provides a further test of how platform-held transaction records enter administrative reporting.", "Normal", False, False, False),
    ("1.1 Research question and objectives", "Heading 2", False, False, False),
    ("One question drives the proposal:", "Normal", False, False, False),
    ("How much economic activity remains invisible when digital platforms are measured through their reported revenue?", "Normal", True, False, True),
    ("Indonesia provides the main empirical setting. Three objectives follow: first, to measure the invisible wedge across available platform histories from 2020 to 2025; second, to explain large movements using disclosed changes in monetization, customer incentives, revenue recognition, and reporting scope; and third, to compare the platform evidence with national e-commerce and payment records to determine whether they produce the same account of economic scale and growth.", "Normal", False, False, False),
    ("A separate comparison of eight platform businesses outside Indonesia tests whether the same transaction-revenue separation appears under other business models.", "Normal", False, False, False),
    ("1.2 Contributions", "Heading 2", False, False, False),
    ("First, the study develops the invisible wedge from a one-year estimate into a measure that can be followed over time. From 2020 to 2025, revenue grows faster than transaction value in ten of twelve annual comparisons, showing that platform revenue is not a stable measure of either the level or the growth of platform-mediated commerce.", "Normal", False, False, False),
    ("Second, I examine why the wedge moves using accounting components disclosed by the platforms themselves. Between 2022 and 2023, Tokopedia's transaction value fell 8.9 percent while selected third-party net revenue rose 53.2 percent. About 60.6 percent of the arithmetic increase in net revenue is associated with lower customer incentives, with the remainder associated with higher gross revenue.", "Normal", False, False, False),
    ("Third, I extend the comparison beyond company accounts. National e-commerce statistics show that almost all of Indonesia's 2023–2024 growth occurred outside the marketplace component, while digital payment measures grew substantially faster than e-commerce value. The records therefore produce different conclusions about the scale, growth, and location of digital activity.", "Normal", False, False, False),
    ("Fourth, I test whether the main findings survive alternative assumptions, sample rules, and external platform evidence rather than depending on one preferred construction.", "Normal", False, False, False),
    ("1.3 Positioning the contribution", "Heading 2", False, False, False),
    ("The contribution is a finance measurement problem before it is a tax-policy one. Platform accounts make the intermediation boundary observable, but that boundary affects the conclusions investors, researchers, statistical agencies, and policymakers draw about scale, growth, and monetization. De Franco, Kothari and Verdi (2011) show why comparable accounting bases matter, while Berg et al. (2020) show that intermediary-held digital traces contain information conventional records can miss. This paper applies those ideas to platform economies and shows that the measure chosen to describe digital activity can change the economic conclusion.", "Normal", False, False, False),
]
for args in ch1:
    insert_before(a2, *args[:2], bold=args[2], italic=args[3], center=args[4])

# Chapter 2: approved working freeze.
clear_between("2. Literature Review and Research Gap", "3. Theoretical Framework and Key Variables")
a3 = find_para("3. Theoretical Framework and Key Variables")
ch2 = [
    ("Digital platforms sit between economic activity and the records used to measure it. They facilitate transactions between buyers and sellers, recognize only part of that activity as their own revenue, and generate records that can also appear in payment, statistical, and administrative systems. Existing research explains each part of this process from a different angle. Taken together, it provides the foundation for asking how much activity remains invisible when platforms are measured through reported revenue.", "Normal"),
    ("2.1 Platform Economics and Revenue Recognition", "Heading 2"),
    ("Platform economics explains why the value of commerce processed through a platform can be much larger than the revenue the platform records. Digital platforms often connect consumers with merchants or other service providers and earn commissions, fees, advertising, logistics, or other intermediation revenue rather than the full value of every transaction they facilitate (Rochet and Tirole 2003; Caillaud and Jullien 2003; Parker and Van Alstyne 2005; Armstrong 2006; Hagiu and Wright 2015; Evans and Schmalensee 2016).", "Normal"),
    ("Accounting can widen or narrow this difference. Under IFRS 15, whether a company acts as a principal or an agent affects whether it recognizes the full customer payment or only the amount it earns from arranging the transaction (International Accounting Standards Board 2014). Customer incentives, monetization, service mix, acquisitions, and changes in business scope can also move reported revenue without an equivalent change in transaction activity. De Franco, Kothari and Verdi (2011) show why comparable accounting bases matter when financial information is interpreted. The literature therefore explains why a transaction–revenue gap can exist and move over time, but not how economically large that gap becomes in practice.", "Normal"),
    ("2.2 Digital Records and Third-Party Information", "Heading 2"),
    ("Activity outside platform revenue is not necessarily unrecorded. Merchants and service providers can be small, self-employed, or weakly represented in conventional financial records while their transactions still leave detailed records inside a digital platform. This differs from the usual shadow-economy problem, where part of the difficulty is that economic activity leaves few reliable formal records (La Porta and Shleifer 2014; Ulyssea 2018; Medina and Schneider 2019). Digital records can contain useful economic information that conventional records miss (Berg et al. 2020), and platform-mediated activity can create additional financial and administrative traces (Barrios, Hochberg and Yi 2022; Denes, Lagaras and Tsoutsoura 2025).", "Normal"),
    ("These records also matter once they become available to government. Public-finance research shows that compliance and enforcement change when information about income or transactions is independently reported by third parties (Kleven et al. 2011; Pomeranz 2015; Naritomi 2019; Kleven, Kreiner and Saez 2016; Slemrod 2019). The OECD Model Rules and the European Union's DAC7 regime apply this principle to digital platforms. Indonesia's Minister of Finance Regulation No. 37 of 2025 follows the same logic through seller identification, transaction-linked reporting, and marketplace withholding. The relevance here is informational: transaction records outside platform revenue may still exist and become usable elsewhere.", "Normal"),
    ("2.3 Measuring the Digital Economy", "Heading 2"),
    ("Official measurement frameworks already distinguish between the commerce facilitated through a platform and the service provided by the platform itself. National-accounting guidance separates the buyer-seller transaction from the platform's intermediation service rather than treating the full transaction value as platform output (Ahmad and Schreyer 2016; International Monetary Fund 2018; OECD 2023; United Nations et al. 2025).", "Normal"),
    ("In practice, however, gross transaction value, platform revenue, e-commerce statistics, and payment data are all used to describe digital economic activity. They measure different things. A difference that is harmless when the measures move together becomes important when they do not, because the choice of measure can change conclusions about economic scale and growth. This thesis therefore compares the records rather than assuming that one can substitute for another.", "Normal"),
    ("2.4 Research Gap", "Heading 2"),
    ("Existing research explains the pieces of the problem separately. Platform economics explains why transaction value can exceed revenue; accounting explains why the gap can move; research on digital records explains why activity outside revenue can still be observable; and official measurement frameworks explain why these quantities should not be treated as equivalent.", "Normal"),
    ("What remains less clear is how large these differences become in practice, whether they persist over time, and whether choosing one measure instead of another changes the economic conclusion. This thesis addresses that gap by measuring the invisible wedge over time, examining disclosed reasons for major movements, comparing platform accounts with national e-commerce and payment records, and using external platform evidence to test whether the same measurement problem appears beyond the main Indonesian setting.", "Normal"),
]
for text, style in ch2:
    insert_before(a3, text, style)

doc.save(PATH)
print(f"updated {PATH.relative_to(ROOT)}")
