# -*- coding: utf-8 -*-
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

R='/home/phyrexian/Downloads/Invisible-ledger/'
OUT='/home/phyrexian/Downloads/llm_automation/project_portfolio/Solarpunk-bitcoin/Invisible_Ledger_Data_Package_2026-09-14.xlsx'
rows=lambda p: list(csv.DictReader(open(R+p)))

wb=Workbook(); wb.remove(wb.active)
HDR=Font(bold=True,size=10,color='FFFFFF'); FILL=PatternFill('solid',fgColor='1F3864')
BOLD=Font(bold=True,size=10); NORM=Font(size=10); TITLE=Font(bold=True,size=12)
THIN=Border(*[Side(style='thin',color='BFBFBF')]*4)

def sheet(name, title, note, header, data, widths=None, freeze='A4'):
    ws=wb.create_sheet(name)
    ws['A1']=title; ws['A1'].font=TITLE
    ws['A2']=note; ws['A2'].font=Font(size=9,italic=True); ws['A2'].alignment=Alignment(wrap_text=True,vertical='top')
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(len(header),4)); ws.row_dimensions[2].height=30
    for j,h in enumerate(header,1):
        c=ws.cell(row=3,column=j,value=h); c.font=HDR; c.fill=FILL; c.border=THIN
        c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,row in enumerate(data,4):
        for j,v in enumerate(row,1):
            c=ws.cell(row=i,column=j,value=v); c.font=NORM; c.border=THIN
            c.alignment=Alignment(wrap_text=True,vertical='top')
    for j,w in enumerate(widths or [18]*len(header),1): ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes=freeze
    return ws

# --- 1 GUIDE ---
g=wb.create_sheet('1_Guide')
g['A1']='Invisible Ledger — empirical data package'; g['A1'].font=Font(bold=True,size=14)
guide=[
 ('Prepared for','Prof. De-Rong Kong (孔德蓉), thesis advisor'),
 ('Prepared by','Christopher Ongko (王新福)'),
 ('Date','14 September 2026'),
 ('Purpose','Summarises the empirical data behind the thesis proposal, so the data can be reviewed before any further manuscript revision.'),
 ('',''),
 ('How to read this file',''),
 ('Sheet 2  FY2023 cross-section','The three-platform Indonesia-focused comparison. This is the main descriptive table.'),
 ('Sheet 3  Source inputs','Every input behind sheet 2, with source file, locator and whether it is directly disclosed or derived.'),
 ('Sheet 4  Candidate inventory','All 17 retained platform-year observations, each labelled with its evidence tier.'),
 ('Sheet 5  Transitions','The 12 within-series annual growth comparisons.'),
 ('Sheet 6  Two universes','Why two valid transition summaries exist (12 and 9) and how they differ.'),
 ('Sheet 7  Sensitivity','One-at-a-time parameter variation and leave-one-platform-out results.'),
 ('Sheet 8  Tokopedia mechanism','The disclosed revenue components behind the FY2022-FY2023 divergence.'),
 ('Sheet 9  BPS official','National e-commerce evidence, 2022-2024.'),
 ('Sheet 10 Global corroboration','48 matched issuer-years outside Indonesia, retained as corroboration only.'),
 ('Sheet 11 Verification ledger','46 claims in the proposal, each re-checked against the source file that produces it.'),
 ('',''),
 ('Three things to note before reading',''),
 ('1. No strictly country-labelled pair exists','No Indonesian platform publishes matched Indonesia transaction value and Indonesia revenue. Every observation is therefore labelled by evidence tier, and tiers are never added together into one monetary total.'),
 ('2. GoTo Group has been dropped','Following your 7 September instruction, GoTo Group figures are no longer used as an Indonesia observation. The Tokopedia e-commerce segment replaces them.'),
 ('3. Nothing here is an approved sample','The inventory is candidate evidence. Which tiers enter the final main sample is a decision for you and the committee; two specific decisions are set out on sheet 6.'),
]
for i,(a,b) in enumerate(guide,3):
    g.cell(row=i,column=1,value=a).font=BOLD if a and not a[0].isdigit() else NORM
    c=g.cell(row=i,column=2,value=b); c.font=NORM; c.alignment=Alignment(wrap_text=True,vertical='top')
g.column_dimensions['A'].width=30; g.column_dimensions['B'].width=96

# --- 2 FY2023 ---
m=rows('data/indonesia_fy2023/fy2023_indonesia_main_rebuilt.csv')
s=rows('data/indonesia_fy2023/fy2023_indonesia_main_summary.csv')[0]
d=[[r['platform'],r['geographic_scope'],float(r['transaction_value_usd_b']),float(r['platform_revenue_usd_b']),
    float(r['transaction_revenue_gap_usd_b']),round(float(r['gap_to_revenue_ratio']),3),
    r['transaction_value_status'],r['revenue_status']] for r in m]
d.append(['TOTAL (three cases)','not an Indonesia-wide total',round(float(s['total_transaction_value_usd_b']),3),
          round(float(s['total_platform_revenue_usd_b']),3),round(float(s['total_transaction_revenue_gap_usd_b']),3),
          round(float(s['aggregate_gap_to_revenue_ratio']),3),'sum of three documented cases','mixed evidence classes'])
sheet('2_FY2023','FY2023 Indonesia-focused cross-section',
 'Values in US$ billions. Wedge = transaction value minus platform revenue. Ecosystem Ratio = wedge / revenue. The total row sums three documented cases under mixed evidence classes; it is not a measurement of the Indonesian platform economy as a whole.',
 ['Platform','Geographic scope','Transaction value V','Platform revenue R','Wedge W = V−R','Ecosystem Ratio E','Transaction value status','Revenue status'],
 d,[26,30,16,16,16,15,30,30])

# --- 3 SOURCE INPUTS ---
si=rows('data/indonesia_fy2023/fy2023_indonesia_source_inputs.csv')
sheet('3_Source_inputs','Source inputs behind the FY2023 cross-section',
 'Every figure used in sheet 2, with its source document and locator. "input_status" states whether the value is directly disclosed by the company, derived from another disclosed figure, or taken from an external market source.',
 ['Input','Platform','Period','Geography','Measure','Value','Unit','Status','Source file','Locator','Scope note'],
 [[r['input_id'],r['platform'],r['period'],r['geography'],r['measure'],r['value'],r['unit'],r['input_status'],
   r['source_file'],r['source_locator'],r['scope_note']] for r in si],
 [14,14,10,16,16,14,12,26,30,30,42])

# --- 4 CANDIDATE INVENTORY ---
lv=rows('outputs/hypothesis_tests_2026-09-10/indonesia_longitudinal_candidate_levels.csv')
sheet('4_Inventory','Candidate platform-year inventory (17 observations)',
 'All retained observations across five issuer series, FY2019–FY2025. Each row carries its evidence tier and the reason its scope is or is not strictly Indonesian. Tiers are never pooled into one monetary total. Tokopedia FY2021 and Bukalapak FY2024 are excluded for period mismatch and do not appear here.',
 ['Series','Year','Transaction value','Revenue','Currency','Unit','Evidence tier','Admission status','Geographic scope','Scope warning'],
 [[r['series'],r['year'],r['transaction_value'],r['revenue_value'],r['currency'],r['unit'],
   r['evidence_tier'],r['admission_status'],r['geography_scope'],r['scope_warning']] for r in lv],
 [26,8,18,16,10,12,30,30,32,46])

# --- 5 TRANSITIONS ---
tr=rows('outputs/hypothesis_tests_2026-09-10/indonesia_longitudinal_candidate_transitions.csv')
sheet('5_Transitions','Within-series annual transitions (12)',
 'Year-on-year growth in transaction value against growth in platform revenue, within each series. A negative "revenue minus transaction" value means transaction value grew faster. Rows where the two move in opposite directions are the cases the thesis examines most closely.',
 ['Series','Transition','Evidence tier','Transaction growth %','Revenue growth %','Revenue − transaction (pp)','Absolute difference (pp)'],
 [[r['series'],r['transition'],r['evidence_tier'],round(float(r['transaction_growth_pct']),2),
   round(float(r['revenue_growth_pct']),2),round(float(r['revenue_minus_transaction_growth_pp']),2),
   round(float(r['absolute_growth_difference_pp']),2)] for r in tr],
 [26,14,30,18,18,22,20])

# --- 6 TWO UNIVERSES + DECISIONS ---
su=rows('outputs/hypothesis_tests_2026-09-10/indonesia_longitudinal_transition_summary.csv')
u=[[r['evidence_tier'],r['transitions'],r['series'],r['revenue_grows_faster'],r['transaction_grows_faster'],
    r['opposite_sign_transitions'],round(float(r['median_absolute_difference_pp']),2)] for r in su]
u.append(['direct_candidate_sensitivity (adds Blibli FY2020, removes Grab/Shopee)','9','3','6','3','3',42.94])
ws=sheet('6_Two_universes','Two valid transition summaries, and the decisions they depend on',
 'The all-tier inventory and the direct-candidate sensitivity are different admissible evidence universes, not a filter applied to one another. The sensitivity removes the conditional Grab and Shopee reconstructions AND adds Blibli FY2020 from the prospectus, which is why it has fewer transitions but one more sign reversal. Both are reported; neither is presented as the approved sample.',
 ['Evidence universe','Transitions','Series','Revenue faster','Transaction faster','Opposite-sign','Median absolute divergence (pp)'],
 u,[52,14,10,16,18,16,22])
r0=len(u)+6
ws.cell(row=r0,column=1,value='Two decisions are requested from the advisor / committee:').font=BOLD
ws.cell(row=r0+1,column=1,value='1. Do Blibli 3P Retail and Bukalapak Group enter the final main sample, remain a labelled scope-pending tier, or are they excluded?').font=NORM
ws.cell(row=r0+2,column=1,value='2. May the Grab and Shopee conditional country reconstructions appear in the main comparison, or should they remain sensitivity evidence only?').font=NORM
ws.cell(row=r0+3,column=1,value='Both decisions change the reported transition counts above.').font=Font(size=10,italic=True)

# --- 7 SENSITIVITY ---
ow=rows('data/indonesia_fy2023/fy2023_indonesia_main_one_way_sensitivity.csv')
gk=[c for c in ow[0] if 'gap' in c and 'usd' in c][0]
sd=[[r['varied_parameter'],r['assumption_class'],r['scenario_value'],r['scenario_unit'],round(float(r[gk]),3)] for r in ow]
lo=rows('data/indonesia_fy2023/fy2023_indonesia_leave_one_platform_out.csv')
lk=[c for c in lo[0] if 'gap' in c and 'usd' in c][0]
for r in lo: sd.append([f"leave-one-out: exclude {r['excluded_platform']}",'composition check',r['remaining_platform_cases']+' cases','platforms',round(float(r[lk]),3)])
sheet('7_Sensitivity','Sensitivity and composition checks',
 'Each assumed parameter is varied one at a time, holding the others fixed, and the selected-platform wedge is recomputed. The leave-one-out rows drop one platform entirely. Baseline wedge is US$40.070 billion.',
 ['Varied parameter / check','Assumption class','Scenario value','Unit','Resulting wedge US$bn'],
 sd,[46,26,22,16,22])

# --- 8 TOKOPEDIA ---
tc=rows('data/measurement/source_extracts/tokopedia_fy2022_2023_revenue_components.csv')
sheet('8_Tokopedia','Tokopedia FY2022–FY2023 revenue components',
 'The disclosed segment components behind the divergence. Gross revenue plus customer incentives equals net revenue in both years. Net revenue rises 53.20% while transaction value falls 8.90%; of that increase, 60.56% is associated with lower customer incentives and 39.44% with higher gross revenue. This is an arithmetic reconciliation of reported figures, not a causal decomposition.',
 ['Period','Component','Value (IDR million)','Source file','Locator','Status'],
 [[r['period'],r['component'],float(r['value_idr_million']),r['source_file'],r['source_locator'],r['source_status']] for r in tc],
 [12,30,20,52,42,30])

# --- 9 BPS ---
nl=rows('outputs/hypothesis_tests_2026-09-10/bps_national_levels.csv')
ga=rows('outputs/hypothesis_tests_2026-09-10/bps_national_growth_anatomy.csv')
bd=[['LEVELS','','','','','']]
for r in nl: bd.append([r['year'],r['transaction_value_idr_trillion'],r['estimated_ecommerce_businesses'],
                        r['implied_idr_million_per_business'],r['marketplace_value_idr_trillion'],r['nonmarketplace_value_idr_trillion']])
bd.append(['GROWTH %','total value','businesses','value per business','marketplace','non-marketplace'])
for r in ga: bd.append([r['transition'],r.get('total_transaction_value_growth_pct',''),r.get('estimated_businesses_growth_pct',''),
                        r.get('implied_value_per_business_growth_pct',''),r.get('marketplace_component_growth_pct',''),r.get('nonmarketplace_component_growth_pct','')])
ws=sheet('9_BPS','BPS official e-commerce evidence, 2022–2024',
 'National indicators from BPS-Statistics Indonesia. Values in IDR trillions. Note the source conflict recorded below: the 2023 business count is taken from the main body/figure because it reconciles to the displayed 2022 count and stated growth rate.',
 ['Year / transition','Transaction value','Businesses','Value per business','Marketplace','Non-marketplace'],
 bd,[20,22,20,22,20,22])
r0=len(bd)+6
ws.cell(row=r0,column=1,value='Source conflict, disclosed:').font=BOLD
ws.cell(row=r0+1,column=1,value='The BPS 2023 publication reports 3,816,750 e-commerce businesses in the main body/figure and 3,934,981 in an executive-summary passage. The former is used because it reconciles to the displayed 2022 count and the stated growth rate. Business-count growth is therefore reported as descriptive context, not as a load-bearing result.').font=NORM
ws.cell(row=r0+2,column=1,value='Province evidence covers 2023 and 2024; the within-province change analysis uses 36 common complete provinces.').font=NORM

# --- 10 GLOBAL ---
gs=rows('data/global_ecommerce/global_platform_issuer_summary.csv')
sheet('10_Global','Global corroboration — 48 matched issuer-years, 8 businesses',
 'Retained as external corroboration only. These are not Indonesian observations, are not pooled with the Indonesia sample, and do not validate the Indonesia country allocations. Take rates range from 0.22% to 74.63%, which is why the wedge is treated as a general property of platform intermediation whose magnitude is business-model determined.',
 ['Issuer','Segment','Region','First year','Last year','Matched years','First take rate %','Last take rate %'],
 [[r['issuer'],r['segment'],r['region'],r['first_year'],r['last_year'],r['matched_years'],
   round(float(r['first_take_rate_pct']),2),round(float(r['last_take_rate_pct']),2)] for r in gs],
 [18,28,30,12,12,14,16,16])

# --- 11 LEDGER ---
lg=rows('outputs/verification_2026-09-14/merged_candidate_claim_ledger.csv')
sheet('11_Verification','Verification ledger — every figure used in the proposal',
 'Each claim in the proposal, the value stated there, the value reproduced from the underlying source file, and the file that produces it. 46 of 46 reproduce.',
 ['Claim','Stated in proposal','Reproduced from source','Status','Source file'],
 [[r['claim'],r['stated_in_proposal'],r['reproduced_value'],r['status'],r['source']] for r in lg],
 [48,18,22,12,52])

wb.save(OUT)
print('saved:',OUT.split('/')[-1])
print('sheets:',len(wb.sheetnames)); [print('   ',n) for n in wb.sheetnames]
