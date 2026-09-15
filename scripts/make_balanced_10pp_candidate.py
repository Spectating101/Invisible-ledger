from pathlib import Path
from docx import Document
from docx.shared import Pt

SOURCE = Path("papers/current/Invisible_Ledger_Thesis_Proposal_FINAL_2026-09-14.docx")
OUT = Path("papers/candidates/Invisible_Ledger_Thesis_Proposal_BALANCED_10PP_CANDIDATE_2026-09-15.docx")


def paragraph_texts(doc):
    return [p.text for p in doc.paragraphs]


def table_texts(doc):
    return [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]


doc = Document(SOURCE)
source_paragraphs = paragraph_texts(doc)
source_tables = table_texts(doc)

section5 = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "5. Preliminary Evidence and Feasibility")
refs = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "References")

# Keep pages 1-5 on the current compact rhythm. From Section 5 onward, relax
# leading and paragraph separation so the empirical / limitations / work-plan
# sequence reads less densely without changing text or shrinking any fonts.
for i, p in enumerate(doc.paragraphs):
    if i < section5:
        continue
    text = p.text.strip()
    fmt = p.paragraph_format
    style = p.style.name

    if i > refs:
        # Restore references to full body size and spread them across the final
        # two pages instead of compressing them into the 9-page tail.
        fmt.line_spacing = 1.25
        fmt.space_after = Pt(6.5)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        continue

    if style.startswith("Heading 1"):
        fmt.space_before = Pt(10)
        fmt.space_after = Pt(4)
    elif style.startswith("Heading 2"):
        fmt.space_before = Pt(7)
        fmt.space_after = Pt(2.5)
    elif text.startswith("Table ") or text.startswith("Table A1."):
        fmt.space_before = Pt(4.5)
        fmt.space_after = Pt(2.5)
        fmt.line_spacing = 1.0
    elif text.startswith("Figure 1."):
        fmt.space_before = Pt(3.5)
        fmt.space_after = Pt(2.5)
        fmt.line_spacing = 1.0
    else:
        fmt.line_spacing = 1.14
        fmt.space_after = Pt(6)
        if text.startswith(("Table 2 reports", "Table 3 compares", "Table 4 summarizes")):
            fmt.space_before = Pt(8)
        elif text.startswith("Figure 1 makes"):
            fmt.space_before = Pt(7)

# Table 1 is already the clean compact table on main. Relax the remaining tables
# only slightly; keep all table text at the existing 9 pt size.
for table in doc.tables[1:]:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.03

assert paragraph_texts(doc) == source_paragraphs
assert table_texts(doc) == source_tables

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

# Re-open and assert the saved candidate is still text-identical.
check = Document(OUT)
assert paragraph_texts(check) == source_paragraphs
assert table_texts(check) == source_tables
print(f"Built {OUT}")
print(f"Content invariant: {len(source_paragraphs)} paragraphs, {len(source_tables)} tables unchanged")
